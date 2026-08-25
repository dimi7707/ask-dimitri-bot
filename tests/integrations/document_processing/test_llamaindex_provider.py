import pytest

from app.integrations.document_processing.base import UnsupportedDocumentTypeError
from app.integrations.document_processing.llamaindex_provider import LlamaIndexDocumentProcessor

PROFILE_TEXT = "Dimitri Avila es desarrollador backend y trabaja con Python, FastAPI y AWS."


def _write_pdf(path, text: str) -> str:
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, text)
    pdf.save()
    return str(path)


def _write_docx(path, text: str) -> str:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    document.save(str(path))
    return str(path)


def _write_pptx(path, text: str) -> str:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = text
    presentation.save(str(path))
    return str(path)


@pytest.fixture
def processor() -> LlamaIndexDocumentProcessor:
    return LlamaIndexDocumentProcessor(chunk_size=512, chunk_overlap=50)


@pytest.mark.parametrize(
    ("writer", "filename"),
    [
        (_write_pdf, "profile.pdf"),
        (_write_docx, "profile.docx"),
        (_write_pptx, "profile.pptx"),
    ],
    ids=["pdf", "docx", "pptx"],
)
def test_load_and_chunk_extracts_text_from_supported_file_types(processor, tmp_path, writer, filename):
    path = writer(tmp_path / filename, PROFILE_TEXT)

    chunks = processor.load_and_chunk(path)

    assert chunks, f"expected at least one chunk for {filename}"
    assert "Dimitri Avila" in " ".join(chunk.chunk_text for chunk in chunks)


def test_load_and_chunk_numbers_chunks_sequentially_from_zero(processor, tmp_path):
    long_text = " ".join(f"Dimitri trabajó en el proyecto número {index}." for index in range(400))
    path = _write_docx(tmp_path / "long.docx", long_text)

    chunks = LlamaIndexDocumentProcessor(chunk_size=64, chunk_overlap=8).load_and_chunk(path)

    assert len(chunks) > 1
    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))


def test_load_and_chunk_records_the_source_file_name_in_metadata(processor, tmp_path):
    path = _write_docx(tmp_path / "profile.docx", PROFILE_TEXT)

    chunks = processor.load_and_chunk(path)

    assert all(chunk.metadata["file_name"] == "profile.docx" for chunk in chunks)


def test_load_and_chunk_rejects_an_unsupported_extension(processor, tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text(PROFILE_TEXT)

    with pytest.raises(UnsupportedDocumentTypeError, match=".txt"):
        processor.load_and_chunk(str(path))


def test_load_and_chunk_rejects_an_unsupported_extension_before_reading_the_file(processor, tmp_path):
    """The extension guard runs first, so a missing/unreadable file still fails with the clear error."""
    with pytest.raises(UnsupportedDocumentTypeError):
        processor.load_and_chunk(str(tmp_path / "does-not-exist.csv"))


def test_load_and_chunk_accepts_uppercase_extensions(processor, tmp_path):
    path = _write_docx(tmp_path / "PROFILE.DOCX", PROFILE_TEXT)

    chunks = processor.load_and_chunk(path)

    assert "Dimitri Avila" in " ".join(chunk.chunk_text for chunk in chunks)
